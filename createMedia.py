import base64
import html, io
import os
import shutil
import sqlite3
import time

import cv2

from dash import html, dcc, Output, Input, State, no_update, ctx, ALL, MATCH
from docx import Document
import dash_daq as daq

#from ClipLocker.logger import log
from app import app
import app as appModule
from styles import *
from embedding import find_best_clip_ids_for_sentences
from tagVideo import renderLeftPreviewAndButton, renderRightDropDowns

VIDEO_EXT = [".mp4", ".mov", ".webm"]
IMAGE_EXT = [".jpg", ".jpeg", ".png", ".webp"]
EXTENSIONS = IMAGE_EXT + VIDEO_EXT


def renderPage3(appSettings):
    return html.Div([
        html.Div(
            id="page-content-wrapper",
            children=[
                html.Div(
                    [
                        # ── Upload ──────────────────────────────────────────────
                        html.Div(
                            dcc.Upload(
                                id="upload-word",
                                multiple=False,
                                children=html.Div("docx hereinziehen"),
                                style={
                                    "width": "100%",
                                    "height": "100%",
                                    "display": "flex",
                                    "alignItems": "center",
                                    "justifyContent": "center",
                                    "textAlign": "center",
                                },
                            ),
                            id = "upload-wrapper",
                            style={
                                "height": "50px",
                                "display": "flex",
                                "alignItems": "center",
                                "justifyContent": "center",
                                "cursor": "pointer",
                                "border": "2px dashed #777",
                                "borderRadius": "10px",
                                "padding": "0 16px",
                                "backgroundColor": "#f8f8f8",
                            },
                        ),

                        # ── Epoche-Dropdown (initial versteckt) ─────────────────
                        html.Div(
                            dcc.Dropdown(
                                id="epoch-dropdown-p3",
                                options=[
                                    {"label": "🏰 Medieval", "value": "medieval"},
                                    {"label": "🏛️ Rome", "value": "rome"},
                                    {"label": "⬜ Epochenlos", "value": "none"},
                                ],
                                placeholder="Epoche wählen...",
                                clearable=False,
                                style={"width": "160px", "color": "black"},
                            ),
                            id="epoch-dropdown-wrapper",
                            style={"display": "none"},
                        ),

                        # ── Bestätigen-Button (initial versteckt) ───────────────
                        html.Div(
                            html.Button(
                                "▶ Bestätigen",
                                id="confirm-script-btn",
                                n_clicks=0,
                                style={
                                    **titleBtnStyle,
                                    "fontSize": "15px",
                                    "cursor": "pointer",
                                    "backgroundColor": "white",
                                    "color": "black",
                                },
                            ),
                            id="confirm-btn-wrapper",
                            style={"display": "none"},
                        ),

                        # ── Titel absolut zentriert ──────────────────────────────
                        html.Div(
                            "ClipLocker",
                            style=headerStyleSmall
                        ),

                        # ── Ordner erstellen ────────────────────────────────────
                        html.Div(
                            html.Button(
                                "Ordner erstellen",
                                id="create-folder",
                                n_clicks=0,
                                style={**titleBtnStyle, "font-size": "15px"}
                            ),
                            style={
                                "display": "flex",
                                "alignItems": "center",
                            },
                        ),
                    ],
                    style={
                        "position": "relative",
                        "display": "flex",
                        "justifyContent": "space-between",
                        "alignItems": "center",
                        "padding": "10px 20px",
                        "borderBottom": "1px solid #ccc",
                        "backgroundColor": "#ffffff",
                        "width": "100%",
                        "boxSizing": "border-box",
                    },
                ),

                # ── Haupt-Inhaltsbereich ─────────────────────────────────────────
                html.Div(
                    [
                        html.Div(
                            id="left-panel",
                            children=[
                                html.Div(id="sentence-panel")
                            ],
                            style={
                                "width": "30%",
                                "padding": "10px",
                                "display": "Block",
                            }
                        ),

                        html.Div(
                            id="right-panel",
                            style={
                                "width": "70%",
                                "padding": "10px",
                                "boxSizing": "border-box",
                                "overflowY": "auto",
                                "height": "80vh",
                                "display": "Block",
                            },
                            children = [
                                html.Div(
                                    [
                                        # =====================================================
                                        # 1. AUSGEWÄHLTE CLIPS
                                        # =====================================================
                                        html.Div(
                                            [
                                                html.Div(
                                                    id="selected-clips-panel",
                                                    style={
                                                        "display": "flex",
                                                        "flexWrap": "wrap",
                                                        "gap": "12px",
                                                        "minHeight": "150px",
                                                        "alignItems": "flex-start",
                                                    }
                                                ),
                                            ],
                                            style={
                                                "border": "1px solid #555",
                                                "borderRadius": "10px",
                                                "padding": "14px",
                                                "marginBottom": "18px",
                                                # "backgroundColor": grey,
                                            }
                                        ),

                                        # =====================================================
                                        # 2. UNTERER BEREICH MIT TABS
                                        # =====================================================
                                        html.Div(
                                            [
                                                # -------------------------
                                                # Tab-Leiste
                                                # -------------------------
                                                html.Div(
                                                    [
                                                        html.Button(
                                                            "Vorschläge",
                                                            id="tab-suggestions",
                                                            n_clicks=0,
                                                            style={
                                                                "padding": "10px 18px",
                                                                "border": "1px solid #666",
                                                                "borderBottom": "none",
                                                                "borderTopLeftRadius": "10px",
                                                                "borderTopRightRadius": "10px",
                                                                "backgroundColor": "#ffffff",
                                                                "color": "black",
                                                                "fontWeight": "bold",
                                                                "cursor": "pointer",
                                                                "marginRight": "4px",
                                                                "position": "relative",
                                                                "top": "1px",
                                                            }
                                                        ),
                                                        html.Button(
                                                            "Clips suchen",
                                                            id="tab-search",
                                                            n_clicks=0,
                                                            style={
                                                                "padding": "10px 18px",
                                                                "border": "1px solid #666",
                                                                "borderBottom": "none",
                                                                "borderTopLeftRadius": "10px",
                                                                "borderTopRightRadius": "10px",
                                                                "backgroundColor": "#d9d9d9",
                                                                "color": "black",
                                                                "fontWeight": "bold",
                                                                "cursor": "pointer",
                                                                "marginRight": "4px",
                                                            }
                                                        ),
                                                        html.Button(
                                                            "Keinen Clip benutzen",
                                                            id="no-clip-btn",
                                                            n_clicks=0,
                                                            style={**titleBtnStyle, "backgroundColor": lightRed,
                                                                   "fontSize": "14px",
                                                                   "marginBottom": "1px"}
                                                        ),
                                                    ],
                                                    style={
                                                        "display": "flex",
                                                        "alignItems": "flex-end",
                                                        "marginBottom": "0",
                                                    }
                                                ),

                                                # -------------------------
                                                # Tab-Inhalt
                                                # -------------------------
                                                html.Div(
                                                    [
                                                        # =====================================
                                                        # TAB 1: Vorschläge
                                                        # =====================================
                                                        html.Div(
                                                            id="suggestions-tab-content",
                                                            children=[
                                                                html.Div(
                                                                    id="video-panel",
                                                                    style={
                                                                        "display": "grid",
                                                                        "gridTemplateColumns": "repeat(auto-fit, minmax(280px, 1fr))",
                                                                        "gap": "12px",
                                                                    },
                                                                ),
                                                            ],
                                                            style={"display": "block"}
                                                        ),

                                                        # =====================================
                                                        # TAB 2: Clips suchen
                                                        # =====================================
                                                        html.Div(
                                                            id="search-tab-content",
                                                            children=[
                                                                html.Div(
                                                                    [
                                                                        dcc.Dropdown(
                                                                            id="search-tags-dropdown",
                                                                            options=getAllTagsFromDB(appSettings),
                                                                            multi=True,
                                                                            placeholder="Tags auswählen...",
                                                                            style={
                                                                                "flex": "1",
                                                                                "color": "black",
                                                                            }
                                                                        ),

                                                                        html.Div(
                                                                            [
                                                                                html.Div(
                                                                                    "OR",
                                                                                    style={
                                                                                        "justifySelf": "end",
                                                                                        "fontWeight": "bold",
                                                                                        "color": "white",
                                                                                        "fontSize": "13px",
                                                                                        "lineHeight": "1",
                                                                                    }
                                                                                ),

                                                                                html.Div(
                                                                                    daq.ToggleSwitch(
                                                                                        id="search-mode-toggle",
                                                                                        value=False,  # False = OR, True = AND
                                                                                        size=28,
                                                                                        color="#22c55e",
                                                                                    ),
                                                                                    style={
                                                                                        "display": "flex",
                                                                                        "justifyContent": "center",
                                                                                        "alignItems": "center",
                                                                                        "marginTop": "-8px",

                                                                                    }
                                                                                ),

                                                                                html.Div(
                                                                                    "AND",
                                                                                    style={
                                                                                        "justifySelf": "start",
                                                                                        "fontWeight": "bold",
                                                                                        "color": "white",
                                                                                        "fontSize": "13px",
                                                                                        "lineHeight": "1",
                                                                                    }
                                                                                ),
                                                                            ],
                                                                            style={
                                                                                "display": "grid",
                                                                                "gridTemplateColumns": "32px 40px 36px",
                                                                                "alignItems": "center",
                                                                                "justifyContent": "center",
                                                                                "columnGap": "4px",
                                                                                "padding": "0 10px",
                                                                                "height": "42px",
                                                                                "border": "1px solid #666",
                                                                                "borderRadius": "999px",
                                                                                "backgroundColor": "#2f2f2f",
                                                                                "boxSizing": "border-box",
                                                                            }
                                                                        ),

                                                                        html.Button(
                                                                            "Suchen",
                                                                            id="search-clips-btn",
                                                                            n_clicks=0,
                                                                            style={
                                                                                "height": "42px",
                                                                                "padding": "0 18px",
                                                                                "border": "1px solid #666",
                                                                                "borderRadius": "8px",
                                                                                "backgroundColor": "#2f2f2f",
                                                                                "color": "white",
                                                                                "fontWeight": "bold",
                                                                                "cursor": "pointer",
                                                                            }
                                                                        ),
                                                                    ],
                                                                    style={
                                                                        "display": "flex",
                                                                        "gap": "10px",
                                                                        "alignItems": "center",
                                                                        "marginBottom": "18px",
                                                                    }
                                                                ),
                                                                html.Div(
                                                                    id="search-results-panel",
                                                                    style={
                                                                        "display": "grid",
                                                                        "gridTemplateColumns": "repeat(auto-fit, minmax(280px, 1fr))",
                                                                        "gap": "20px",
                                                                        "marginBottom": "20px",
                                                                        "alignItems": "start",
                                                                    }
                                                                ),

                                                                html.Div(
                                                                    id="search-pagination-container",
                                                                    children=[
                                                                        html.Button(
                                                                            "← Zurück",
                                                                            id="search-prev-page",
                                                                            n_clicks=0,
                                                                            style={
                                                                                "padding": "8px 14px",
                                                                                "border": "1px solid #666",
                                                                                "borderRadius": "8px",
                                                                                "backgroundColor": "#2f2f2f",
                                                                                "color": "white",
                                                                                "cursor": "pointer",
                                                                            }
                                                                        ),

                                                                        html.Div(
                                                                            id="search-page-info",
                                                                            style={
                                                                                "minWidth": "120px",
                                                                                "textAlign": "center",
                                                                                "fontWeight": "bold",
                                                                            }
                                                                        ),

                                                                        html.Button(
                                                                            "Weiter →",
                                                                            id="search-next-page",
                                                                            n_clicks=0,
                                                                            style={
                                                                                "padding": "8px 14px",
                                                                                "border": "1px solid #666",
                                                                                "borderRadius": "8px",
                                                                                "backgroundColor": "#2f2f2f",
                                                                                "color": "white",
                                                                                "cursor": "pointer",
                                                                            }
                                                                        ),
                                                                    ],
                                                                    style=paginationContainerStyle
                                                                ),
                                                            ],
                                                            style={"display": "none"}
                                                        )
                                                    ],
                                                    style={
                                                        "border": "1px solid #666",
                                                        "borderRadius": "0 10px 10px 10px",
                                                        "padding": "16px",
                                                        # "backgroundColor": "#1f1f1f",
                                                        "minHeight": "420px",
                                                    }
                                                )
                                            ]
                                        )
                                    ]
                                )
                            ]
                        ),
                    ],
                    style={
                        "display": "flex",
                        "flexDirection": "row",
                        "width": "100%",
                    }
                )
            ]
        ),
        html.Div(
            id="tag-edit-overlay",
            children=[
                html.Div(
                    id="tag-edit-modal",
                    children=[
                        html.Div(id="tag-edit-modal-content")
                    ],
                    style=tagEditModalStyle
                )
            ],
            style={"display": "none"}
        )
    ])

@app.callback(
    Output("confirm-script-btn", "disabled"),
    Output("confirm-script-btn", "style"),
    Input("sentence-store", "data"),
    Input("epoch-dropdown-p3", "value"),
    prevent_initial_call=True,
)
def activateConfirmBtn(sentences, epoch):
    if sentences and epoch:
        return False, {**titleBtnStyle, "fontSize": "15px",
                       "cursor": "pointer", "backgroundColor": "white", "color": "black"}
    return True, {**titleBtnStyle, "fontSize": "15px"}


@app.callback(
    Output("epoch-dropdown-wrapper", "style"),
    Output("upload-wrapper", "style"),
    Input("sentence-store", "data"),
    prevent_initial_call=True,
)
def showEpochDropdown(sentences):
    if sentences:
        return {"display": "block"}, {"display": "none"}
    return {"display": "none"}, no_update


@app.callback(
    Output("confirm-btn-wrapper", "style"),
    Output("epoch-dropdown-wrapper", "style", allow_duplicate=True),
    Input("epoch-dropdown-p3", "value"),
    prevent_initial_call=True,
)
def showConfirmBtn(epoch):
    if epoch:
        return {"display": "block"}, {"display": "none"}
    return {"display": "none"}, no_update


@app.callback(
    Output("sentence-store", "data"),
    Output("documentName-store", "data"),
    Input("upload-word", "contents"),
    Input("upload-word", "filename"),
)
def safeSentences(contents, filename):

    if contents is None or filename is None:
        return no_update

    if not filename.endswith(".docx"):
        return no_update

    content_type, content_string = contents.split(",",1)    #aufteilung in header und data(1 stellt sicher nur 1 Teilung)
    decoded = base64.b64decode(content_string)              #decodierung der daten
    file=io.BytesIO(decoded)                                #umwandlung zu datei
    document = Document(file)                               #umwandlunng in dateistruktur
    sentences = [p.text for p in document.paragraphs if p.text.strip()]       #liste aller absätze(Sätze)

    output = []

    for sentence in sentences:
        output.append(sentence)

    print(filename)
    return output, filename

@app.callback(
    Output("sentence-panel", "children"),
    Input("sentence-store", "data"),
    State("app-settings-store", "data"),
    prevent_initial_call=True,
)
def renderSentences(sentences, appSettings):
    if not sentences:
        return []

    containers = []

    for i, sentence in enumerate(sentences):
        containers.append(
            html.Div(
                [
                    html.Button(
                        sentence,
                        id={"type": "sentence_button", "sentence_index": i},
                        n_clicks=0,
                        style=baseSentenceStyle
                    )
                ]
            )
        )

    return containers

@app.callback(
    Output("selectedSentence-store", "data"),
    Input({"type": "sentence_button", "sentence_index": ALL}, "n_clicks"),
    prevent_initial_call=True
)
def saveSelectedSentence(click):
    triggered = ctx.triggered_id

    if not triggered:
        return no_update

    print("Ausgewählter Satz:", triggered["sentence_index"])

    return triggered["sentence_index"]

@app.callback(
    Output({"type": "sentence_button", "sentence_index": ALL}, "style"),
    Input("selectedSentence-store", "data"),
    State("sentence-store", "data"),
    Input("selectedClips-store", "data"),
    Input("placeholder-active-store", "data"),
)
def updateAndDisplaySentenceStates(selectedSentence, sentences, selectedClips, placeHolderActiveStore):

    if not sentences:
        return no_update

    if not isinstance(selectedClips, list):
        selectedClips = []

    if not isinstance(placeHolderActiveStore, list) or len(placeHolderActiveStore) != len(sentences):
        placeHolderActiveStore = [False] * len(sentences)

    styles = []

    for i, sentence in enumerate(sentences):
        style = baseSentenceStyle.copy()

        clipsForSentence = selectedClips[i] if i < len(selectedClips) else []
        hasClipsSelected = len(clipsForSentence) > 0
        placeHolderActive = placeHolderActiveStore[i]

        if placeHolderActive:
            if i == selectedSentence:
                style["border"] = "4px solid red"
            else:
                style["border"] = "2px solid red"

        # 2. Danach normale Clip-Auswahl
        elif hasClipsSelected:
            if i == selectedSentence:
                style["border"] = "4px solid #ADFF2F"
            else:
                style["border"] = "2px solid #ADFF2F"

        elif i == selectedSentence:
            style["border"] = "4px solid black"

        styles.append(style)

    return styles


#--------------------------------------------------------------------------------------------------


@app.callback(
    Output("clips-per-sentence-store", "data"),
    Input("confirm-script-btn", "n_clicks"),
    State("sentence-store", "data"),
    State("epoch-dropdown-p3", "value"),
    State("app-settings-store", "data"),
    prevent_initial_call=True
)
def pickClips(n_clicks, sentences, epoch, appSettings):
    if not sentences or not epoch or not n_clicks:
        return no_update

    print("Picking Clips")
    output=find_best_clip_ids_for_sentences(sentences, appSettings["path"] + appSettings["db"], 20, epoch)
    print(output)
    return output

@app.callback(
    Output("suggestions-tab-content", "style"),
    Output("search-tab-content", "style"),
    Output("tab-suggestions", "style"),
    Output("tab-search", "style"),
    Input("active-right-tab", "data"),
)
def updateRightTabContent(activeTab):
    if activeTab == "search":
        return (
            {"display": "none"},
            {"display": "block"},
            inactiveTabStyle,
            activeTabStyle,
        )

    return (
        {"display": "block"},
        {"display": "none"},
        activeTabStyle,
        inactiveTabStyle,
    )


@app.callback(
    Output("active-right-tab", "data"),
    Input("tab-suggestions", "n_clicks"),
    Input("tab-search", "n_clicks"),
    prevent_initial_call=True,
)
def switchRightTab(nSuggestions, nSearch):
    trigger = ctx.triggered_id

    if trigger == "tab-suggestions":
        return "suggestions"
    if trigger == "tab-search":
        return "search"

    return no_update

def getAllTagsFromDB(appSettings):
    dbPath = appSettings["path"] + appSettings["db"]

    conn = sqlite3.connect(dbPath)
    database = conn.cursor()
    database.execute("SELECT tag_name FROM Tags")
    rows = database.fetchall()
    conn.close()

    return [{"label": row[0], "value": row[0]} for row in rows]


@app.callback(
    Output("search-results-store", "data"),
    Input("search-clips-btn", "n_clicks"),
    State("search-tags-dropdown", "value"),
    State("search-mode-toggle", "value"),
    State("app-settings-store", "data"),
    State("epoch-dropdown-p3", "value"),
    prevent_initial_call=True
)
def saveSearchedClips(btn, selectedTags, searchMode, appSettings, epoch):
    if not selectedTags or not btn:
        return no_update

    dbPath = appSettings["path"] + appSettings["db"]
    placeholders = ",".join("?" * len(selectedTags))

    epochFilter = "AND (c.epoch = ? OR c.epoch IS NULL)" if epoch and epoch != "none" else ""

    conn = sqlite3.connect(dbPath)
    database = conn.cursor()

    if searchMode:  # AND
        query = f"""
              SELECT ct.clip_id
              FROM ClipTag AS ct
              JOIN Tags AS t ON t.tag_id = ct.tag_id
              JOIN Clips AS c ON c.clip_id = ct.clip_id
              WHERE t.tag_name IN ({placeholders})
              {epochFilter}
              GROUP BY ct.clip_id
              HAVING COUNT(DISTINCT t.tag_name) = ?
              ORDER BY ct.clip_id
          """
        params = [*selectedTags, epoch, len(selectedTags)] if epochFilter else [*selectedTags, len(selectedTags)]
    else:  # OR
        query = f"""
              SELECT DISTINCT ct.clip_id
              FROM ClipTag AS ct
              JOIN Tags AS t ON t.tag_id = ct.tag_id
              JOIN Clips AS c ON c.clip_id = ct.clip_id
              WHERE t.tag_name IN ({placeholders})
              {epochFilter}
              ORDER BY ct.clip_id
          """
        params = [*selectedTags, epoch] if epochFilter else selectedTags

    database.execute(query, params)
    rows = database.fetchall()
    conn.close()

    result = [row[0] for row in rows]
    print(result)
    return result


@app.callback(
    Output("search-page-store", "data"),
    Input("search-prev-page", "n_clicks"),
    Input("search-next-page", "n_clicks"),
    Input("search-results-store", "data"),
    State("search-page-store", "data"),
    prevent_initial_call=True
)
def changeSearchPage(prevClicks, nextClicks, searchResults, currentPage):
    if not searchResults:
        return 0

    currentPage = currentPage or 0
    pageSize = 10
    totalPages = (len(searchResults) - 1) // pageSize + 1

    trigger = ctx.triggered_id

    if trigger == "search-results-store":
        return 0

    if trigger == "search-prev-page":
        return max(0, currentPage - 1)

    if trigger == "search-next-page":
        return min(totalPages - 1, currentPage + 1)

    return currentPage


@app.callback(
    Output("search-pagination-container", "style"),
    Input("search-results-store", "data"),
    prevent_initial_call=True
)
def togglePagination(results):
    if not results:
        return { **paginationContainerStyle, "display": "none"}

    return { **paginationContainerStyle, "display": "flex"}


@app.callback(
    Output("search-results-panel", "children"),
    Output("search-page-info", "children"),
    Output("search-prev-page", "disabled"),
    Output("search-next-page", "disabled"),
    Output("search-prev-page", "style"),
    Output("search-next-page", "style"),
    Input("search-results-store", "data"),
    Input("search-page-store", "data"),
    State("app-settings-store", "data"),
    State("selectedClips-store", "data"),
    State("selectedSentence-store", "data"),
    Input("clip-usage-store", "data"),
    prevent_initial_call=True
)
def renderSearchedClips(searchResults, currentPage, appSettings, selectedClips, selectedSentence, clipUsageStore):
    if not searchResults:
        return [], "Seite 0 / 0", True, True, disabledPagingBtnStyle, disabledPagingBtnStyle

    appModule.rawMediaFolder = appSettings["path"] + appSettings["fs"]

    currentPage = currentPage or 0
    pageSize = 20
    totalResults = len(searchResults)
    totalPages = (totalResults - 1) // pageSize + 1

    start = currentPage * pageSize
    end = start + pageSize
    visibleClipIds = searchResults[start:end]

    selectedClips = selectedClips or []
    selectedForThisSentence = []

    if selectedSentence is not None and selectedSentence < len(selectedClips):
        selectedForThisSentence = selectedClips[selectedSentence] or []

    output = []

    for clipId in visibleClipIds:
        filePath = None
        extension = None

        currentClipStyle = clipStyle.copy()


        for ext in EXTENSIONS:
            testPath = os.path.join(appSettings["path"] + appSettings["fs"], f"{clipId}{ext}")
            if os.path.exists(testPath):
                filePath = f"/media/{clipId}{ext}"
                extension = ext
                break



        #if clipId in selectedForThisSentence:
        #    currentClipStyle["border"] = "2px solid #ADFF2F"


        if filePath is None:
            media = html.Div(
                f"Clip {clipId} nicht gefunden",
                style={
                    "width": "300px",
                    "height": "170px",
                    "display": "flex",
                    "alignItems": "center",
                    "justifyContent": "center",
                    "border": "2px solid black",
                    "borderRadius": "8px",
                    "textAlign": "center",
                    "padding": "8px",
                    "backgroundColor": "#f0f0f0",
                    "color": "black",
                }
            )
        elif extension in VIDEO_EXT:
            media = html.Video(
                id={"type": "video", "form": "searched","clip_id": clipId},
                src=filePath,
                controls=True,
                preload="metadata",
                #poster=create_video_thumbnail(filePath),
                style=clipStyle
            )
        elif extension in IMAGE_EXT:
            media = html.Img(
                id={"type": "video","form": "searched", "clip_id": clipId},
                src=filePath,
                style=clipStyle
            )
        else:
            media = html.Div(
                f"Unbekannter Dateityp: {clipId}",
                style={
                    "width": "300px",
                    "height": "170px",
                    "display": "flex",
                    "alignItems": "center",
                    "justifyContent": "center",
                    "border": "2px solid black",
                    "borderRadius": "8px",
                    "textAlign": "center",
                    "padding": "8px",
                    "backgroundColor": "#f0f0f0",
                    "color": "black",
                }
            )

        infoElement = build_clip_info_element(clipId, selectedSentence, clipUsageStore)

        output.append(
            html.Div(
                [
                    html.Div(
                        [
                            media,
                            html.Button(
                                "✓",
                                id={"type": "video_button", "form": "recommended", "clip_id": clipId},
                                n_clicks=0,
                                style=selectClipBtnStyle
                            ),
                            html.Button(
                                "✎",
                                id={"type": "edit_tags_button", "form": "recommended", "clip_id": clipId},
                                n_clicks=0,
                                style={**selectClipBtnStyle, "left": "3px"}
                            ),
                        ],
                        style=clipFrameStyle
                    ),
                    infoElement
                ],
                style=clipPlusBtnContainerStyle
            )
        )

    pageInfo = f"Seite {currentPage + 1} / {totalPages}"

    prevDisabled = currentPage == 0
    nextDisabled = currentPage >= totalPages - 1

    prevStyle = disabledPagingBtnStyle if prevDisabled else basePagingBtnStyle
    nextStyle = disabledPagingBtnStyle if nextDisabled else basePagingBtnStyle

    return output, pageInfo, prevDisabled, nextDisabled, prevStyle, nextStyle


@app.callback(
    Output("video-panel", "children"),
    Input("selectedSentence-store", "data"),
    Input("clips-per-sentence-store", "data"),
    State("selectedClips-store", "data"),
    State("app-settings-store", "data"),
    Input("clip-usage-store", "data"),
    prevent_initial_call=True
)
def renderRecommendedClips(selectedSentence, clips, selectedClips, appSettings, clipUsageStore):
    if selectedSentence is None:
        return no_update

    if not clips:
        return []

    if not appSettings or not appSettings["path"] + appSettings["fs"]:
        return []

    appModule.rawMediaFolder = appSettings["path"] + appSettings["fs"]

    selectedClips = selectedClips or []

    if selectedSentence >= len(clips):
        #log("Keine Clips gefunden für Satz:", selectedSentence, color = "red")
        return []

    output = []

    clipsForThisSentence = clips[selectedSentence]

    for j, clip in enumerate(clipsForThisSentence):

        newVideoStyle = clipStyle.copy()

        # am anfang selected clips leer, daher überprüfung
        #if selectedClips and clip in selectedClips[selectedSentence]:
        #    newVideoStyle["border"] = "3px solid #ADFF2F"

        file_path = None
        extension = None

        for ext in EXTENSIONS:
            test_path = os.path.join(appSettings["path"] + appSettings["fs"], f"{clip}{ext}")

            if os.path.exists(test_path):
                file_path = f"/media/{clip}{ext}"
                extension = ext
                break

        if extension in VIDEO_EXT:
            media = html.Video(
                id={"type": "video","form": "recommended", "clip_id": clip},
                src=file_path,
                controls=True,
                preload="metadata",
                style=newVideoStyle
            )

        elif extension in IMAGE_EXT:
            media = html.Img(
                id={"type": "video","form": "recommended", "clip_id": clip},
                src=file_path,
                style=newVideoStyle
            )

        else:
            media = html.Div("Unbekannter Dateityp")

        infoElement = build_clip_info_element(clip, selectedSentence, clipUsageStore)

        output.append(
            html.Div(
                [
                    html.Div(
                        [
                            media,
                            html.Button("✓",
                                        id={"type": "video_button","form": "recommended", "clip_id": clip},
                                        style=selectClipBtnStyle),
                            html.Button("✎",
                                        id={"type": "edit_tags_button", "form": "recommended", "clip_id": clip},
                                        style={**selectClipBtnStyle, "left": "3px"}),
                        ],
                        style=clipFrameStyle
                    ),
                    infoElement
                ],
                style=clipPlusBtnContainerStyle
            )
        )

    return output

def build_clip_info_element(clip, selectedSentence, clipUsageStore):
    usage = clipUsageStore or {}
    info = usage.get(str(clip), {})

    count = info.get("count", 0)
    last = info.get("last_sentence")
    previous_last = info.get("previous_last_sentence")

    display_last = last
    if last == selectedSentence and previous_last is not None:
        display_last = previous_last

    if count > 0:
        if count == 2:
            countColor = "orange"
        elif count > 2:
            countColor = "red"
        else:
            countColor = "grey"

        prefix = f"#{clip}  ·  "
        count_text = f"{count}x"

        suffix = ""
        if display_last is not None:
            sentences_ago = selectedSentence - display_last

            if sentences_ago > 1:
                suffix = f"  ·  vor {sentences_ago} Sätzen"
            elif sentences_ago == 1:
                suffix = f"  ·  vor 1 Satz"

        return html.Div(
            [
                html.Span(prefix),
                html.Span(count_text, style={"color": countColor}),
                html.Span(suffix),
            ],
            style={
                "fontSize": "14px",
                "color": "#888",
                "textAlign": "center",
                "marginTop": "4px"
            }
        )

    return html.Div(
        f"#{clip}",
        style={
            "fontSize": "14px",
            "color": "#888",
            "textAlign": "center",
            "marginTop": "4px"
        }
    )


@app.callback(
    Output("selectedClips-store", "data", allow_duplicate=True),
    Output("clip-usage-store", "data"),
    Input({"type": "video_button", "form": ALL, "clip_id": ALL}, "n_clicks"),
    State("selectedSentence-store", "data"),
    State("selectedClips-store", "data"),
    State("clips-per-sentence-store", "data"),
    State("placeholder-active-store", "data"),
    State("clip-usage-store", "data"),
    prevent_initial_call=True
)
def saveSelectedClips(clicks, selectedSentence, selectedClips, clips, placeholderActiveStore, clipUsageStore):

    triggered = ctx.triggered_id

    triggered_value = ctx.triggered[0]["value"]

    if not clips:
        return no_update

    if triggered_value in (None, 0):
        return no_update

    if not clicks or all(c in (None, 0) for c in clicks): #wenn clicks überall 0
        return no_update

    if triggered is None:
        return no_update

    if selectedSentence is None:
        return no_update

    if not isinstance(selectedClips, list) or len(selectedClips) != len(clips):
        selectedClips = [[] for _ in range(len(clips))]
    else:
        selectedClips = [list(x) for x in selectedClips]

    clipID = triggered["clip_id"] # zb 2

    if not isinstance(placeholderActiveStore, list) or len(placeholderActiveStore) <= selectedSentence:
        placeholderActive = False
    else:
        placeholderActive = placeholderActiveStore[selectedSentence]

    if clipID not in selectedClips[selectedSentence] and not placeholderActive:
        selectedClips[selectedSentence].append(clipID)

        usage = clipUsageStore or {}
        key = str(clipID)
        old_info = usage.get(key, {})
        old_last_sentence = old_info.get("last_sentence")

        usage[key] = {
            "count": usage.get(key, {}).get("count", 0) + 1,
            "last_sentence": selectedSentence,
            "previous_last_sentence": old_last_sentence,
        }
        return selectedClips, usage

    print("Clips selected: ",selectedClips)
    return selectedClips, no_update


@app.callback(
    #Output({"type": "video","form": "recommended", "clip_id": ALL}, "style"),
    Output("selected-clips-panel", "children"),
    Input("selectedClips-store", "data"),
    Input("selectedSentence-store", "data"),
    Input("placeholder-active-store", "data"),
    State("clips-per-sentence-store", "data"),
    State("app-settings-store", "data"),
    State("clip-usage-store", "data"),
    prevent_initial_call=True
)
def displaySelectedClips(selectedClips, selectedSentence, placeholderStore, clips, appSettings, clipUsageStore):
    if selectedSentence is None:
        return []

    if placeholderStore and selectedSentence < len(placeholderStore):
        if placeholderStore[selectedSentence]:
            return [
                html.Div(
                    html.Img(
                        src="/assets/placeholder.png",
                        style=miniClipStyle
                    )
                )
            ]

    if selectedSentence is None:
        return no_update

    if not clips:
        return []

    if selectedSentence >= len(clips):
        return []

    selectedClips = selectedClips or []

    clipsForThisSentence = clips[selectedSentence]

    if selectedSentence < len(selectedClips):
        selectedClipsForThisSentence = selectedClips[selectedSentence] or []
    else:
        selectedClipsForThisSentence = []

    # ---------------------------------
    # 1. Styles für Vorschlagsclips unten
    # ---------------------------------
    #styles = []


    #for clip in clipsForThisSentence:
     #   style = clipStyle.copy()
      #  if clip in selectedClipsForThisSentence:
       #     style["border"] = "3px solid #ADFF2F"
        #styles.append(style)


    # ---------------------------------
    # 2. Oben ausgewählte Clips rendern
    # ---------------------------------
    selectedClipCards = []

    for clip in selectedClipsForThisSentence:
        file_path = None
        extension = None

        for ext in EXTENSIONS:
            test_path = os.path.join(appSettings["path"] + appSettings["fs"], f"{clip}{ext}")
            if os.path.exists(test_path):
                file_path = f"/media/{clip}{ext}"
                extension = ext
                break

        if file_path is None:
            media = html.Div(
                f"Clip {clip} nicht gefunden",
                style=miniClipErrorStyle
            )
        elif extension in VIDEO_EXT:


            media = html.Video(
                src=file_path,
                controls=False,
                preload="metadata",
                style=miniClipStyle
            )
        elif extension in IMAGE_EXT:
            media = html.Img(
                src=file_path,
                style=miniClipStyle
            )
        else:
            media = html.Div(
                "Unbekannter Dateityp",
                style=miniClipErrorStyle
            )

        infoElement = build_clip_info_element(clip, selectedSentence, clipUsageStore)

        selectedClipCards.append(
            html.Div([
                html.Div([
                    media,
                    html.Button(
                        "🞬",
                        id={"type": "video_cancel_button", "clip_id": clip},
                        n_clicks=0,
                        style={**selectClipBtnStyle, "height": "25px", "width": "25px"}),
                ],
                    style={
                        "position": "relative",
                        "width": "200px",
                        "height": "120px",
                        "flexShrink": "0",
                        "borderRadius": "8px",
                        "overflow": "hidden",
                        "border": "1px solid black",
                    }),
                infoElement
            ],
                style={"flexShrink": "0"})
        )

    return selectedClipCards


def create_video_thumbnail(video_path, time_ms=100):
    print(video_path)
    cap = cv2.VideoCapture(video_path)

    if not cap.isOpened():
        raise RuntimeError("Video konnte nicht geöffnet werden")

    # zu gewünschter Stelle springen (z. B. 0.3 Sekunden)
    cap.set(cv2.CAP_PROP_POS_MSEC, time_ms)

    success, frame = cap.read()
    cap.release()

    if not success:
        raise RuntimeError("Konnte keinen Frame lesen")

    # Frame → JPEG im Speicher encodieren
    success, buffer = cv2.imencode(".jpg", frame)
    if not success:
        raise RuntimeError("Encoding fehlgeschlagen")

    # Base64 erzeugen
    jpg_as_text = base64.b64encode(buffer).decode("utf-8")

    return f"data:image/jpeg;base64,{jpg_as_text}"


@app.callback(
    Output("selectedClips-store", "data", allow_duplicate=True),
    Output("clip-usage-store", "data", allow_duplicate=True),
    Input({"type": "video_cancel_button", "clip_id": ALL}, "n_clicks"),
    State("selectedSentence-store", "data"),
    State("selectedClips-store", "data"),
    State("clip-usage-store", "data"),
    prevent_initial_call=True
)
def removeSelectedClip(cancelClicks, selectedSentence, selectedClips, clipUsageStore):

    triggered = ctx.triggered_id

    if not triggered:
        return no_update

    triggered_value = ctx.triggered[0]["value"]
    if triggered_value in (None, 0):
        return no_update

    if selectedSentence is None or not selectedClips:
        return no_update

    selectedClips = [list(x) for x in selectedClips]
    usage = dict(clipUsageStore or {})

    clipID = triggered["clip_id"]

    if selectedSentence < len(selectedClips) and clipID in selectedClips[selectedSentence]:
        selectedClips[selectedSentence].remove(clipID)
        print("Clip entfernt:", clipID)

        # usage für genau diesen Clip komplett neu berechnen
        occurrences = []

        for sentence_index, clips_in_sentence in enumerate(selectedClips):
            if clipID in clips_in_sentence:
                occurrences.append(sentence_index)

        key = str(clipID)

        if not occurrences:
            # Clip kommt nirgends mehr vor -> ganz aus usage entfernen
            usage.pop(key, None)
        else:
            usage[key] = {
                "count": len(occurrences),
                "last_sentence": occurrences[-1],
                "previous_last_sentence": occurrences[-2] if len(occurrences) >= 2 else None
            }

    print("Clips selected:", selectedClips)
    print("Usage:", usage)
    return selectedClips, usage


@app.callback(
    Output("placeholder-active-store", "data"),
    Input("no-clip-btn", "n_clicks"),
    State("placeholder-active-store", "data"),
    State("selectedSentence-store", "data"),
    State("sentence-store", "data"),
    State("selectedClips-store", "data"),
    prevent_initial_call=True
)
def togglePlaceholder(n_clicks, placeholderStore, selectedSentence, sentences, selectedClips):

    if selectedSentence is None or not sentences:
        return no_update

    if len(selectedClips) != 0:
        if len(selectedClips[selectedSentence]) != 0:
            return no_update


    # initialisieren falls leer
    if not placeholderStore or len(placeholderStore) != len(sentences):
        placeholderStore = [False] * len(sentences)
    else:
        placeholderStore = list(placeholderStore)

    # toggle nur für diesen Satz
    placeholderStore[selectedSentence] = not placeholderStore[selectedSentence]
    print("placeholderStore:", placeholderStore)

    return placeholderStore


@app.callback(
    Output("selectedSentence-store", "data", allow_duplicate=True),#DUMMY
    Input("create-folder", "n_clicks"),
    State("sentence-store", "data"),
    State("selectedClips-store", "data"),
    State("documentName-store", "data"),
    State("app-settings-store", "data"),
    State("placeholder-active-store", "data"),
    prevent_initial_call=True
)
def createFolder(n_clicks, sentences, selectedClips, filename, appSettings, placeholderStore):

    if not selectedClips:
        return no_update

    if not isinstance(placeholderStore, list) or len(placeholderStore) != len(sentences):
        placeholderStore = [False] * len(sentences)

    if not isinstance(selectedClips, list):
        selectedClips = [[] for _ in range(len(sentences))]


    placeholderSource = os.path.join(os.path.dirname(__file__), "assets", "placeholder.png")

    folder = filename.replace(".docx", "")
    folderpath = r"C:\Users\Admin\Documents\05-No_Mans_Sky\Database\\" + folder +"_Media"

    os.makedirs(folderpath, exist_ok=True)

    for i, sentence in enumerate(sentences):
        clipsPerSentence = selectedClips[i] if i < len(selectedClips) else []
        placeholderActive = placeholderStore[i]

        sentence = sentences[i]
        beginning = safe_filename_part(" ".join(sentence.split()[:3]))
        number = str(i + 1).zfill(3)

        if placeholderActive:
            if os.path.exists(placeholderSource):
                destination = os.path.join(
                    folderpath,
                    f"{number}...{beginning}.png"
                )
                shutil.copy(placeholderSource, destination)
            else:
                print(f"Placeholder nicht gefunden: {placeholderSource}")
            continue



        if not clipsPerSentence:
            continue

        for clip in clipsPerSentence:

            source = None
            extension = None

            for ext in EXTENSIONS:
                test_path = os.path.join(appSettings["path"] + appSettings["fs"], f"{clip}{ext}")
                if os.path.exists(test_path):
                    source = test_path
                    extension = ext
                    break

            if source is None:
                print(f"Keine Datei gefunden für Clip {clip}")
                continue

            number = str(i + 1).zfill(3)
            destination = os.path.join(folderpath, f"{number}_({clip})...{beginning}{extension}")

            shutil.copy(source, destination)

    return no_update

def safe_filename_part(text):
    forbidden = '<>:"/\\|?*'
    for ch in forbidden:
        text = text.replace(ch, "")
    return text.strip()


@app.callback(
    Output("edit-tags-store", "data"),
    Input({"type": "edit_tags_button", "form": ALL, "clip_id": ALL}, "n_clicks_timestamp"),
    prevent_initial_call=True
)
def openTagEditor(timestamps):
    triggered = ctx.triggered_id
    if not triggered:
        return no_update

    timestamp = ctx.triggered[0]["value"]
    if not timestamp:
        return no_update

    # Ignorieren wenn Klick älter als 1 Sekunde
    if (time.time() * 1000 - timestamp) > 1000:
        return no_update

    return {
        "clip_id": triggered["clip_id"],
        "form": triggered["form"]
    }


@app.callback(
    Output("tag-edit-overlay", "style"),
    Output("page-content-wrapper", "style"),
    Input("edit-tags-store", "data"),
    prevent_initial_call=False
)
def toggleTagEditorOverlay(editData):
    if editData is None:
        return {"display": "none"}, pageContentNormalStyle

    return tagEditOverlayStyle, pageContentBlurredStyle


@app.callback(
    Output("tag-edit-modal-content", "children"),
    Output("tag-edit-overlay", "style", allow_duplicate=True),
    Output("edit-clip-id-store", "data"),
    Input({"type": "edit_tags_button", "form": ALL, "clip_id": ALL}, "n_clicks"),
    State("app-settings-store", "data"),
    prevent_initial_call=True,
)
def openEditModal(n_clicks, appSettings):
    if not any(n for n in n_clicks if n):
        return no_update, no_update, no_update

    clipId = ctx.triggered_id["clip_id"]

    clip = None
    for ext in EXTENSIONS:
        path = os.path.join(appSettings["path"] + appSettings["fs"], f"{clipId}{ext}")
        if os.path.exists(path):
            clip = {"name": str(clipId), "extension": ext,
                    "type": "video" if ext in VIDEO_EXT else "image"}
            break

    if clip is None:
        return no_update, no_update, no_update

    content = html.Div([
        html.Div(
            style={"display": "flex", "gap": "20px"},
            children=[
                renderLeftPreviewAndButton(clip, editMode=True),
                renderRightDropDowns(
                    appSettings,
                    editMode=True,        # ← kein claudeData
                    editClipId=clipId     # ← aus DB laden
                )
            ]
        )
    ])
    print("clipId: ", clipId)

    return content, tagEditOverlayStyle, clipId


@app.callback(
    Output("tag-edit-overlay", "style", allow_duplicate=True),
    Output("page-content-wrapper", "style", allow_duplicate=True),
    Input("db-button", "n_clicks"),   #nach Speichern schließen
    prevent_initial_call=True,
)
def closeEditModal(btn):
    if btn is None or btn == 0:
        return no_update, no_update

    return {"display": "none"}, pageContentNormalStyle